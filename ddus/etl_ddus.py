"""Лист «Бюджет ДДУС» -> плоская таблица ddus_budget (CSV и/или PostgreSQL).

Запуск:
    python etl_ddus.py <файл.xlsx> --csv ddus_budget.csv
    python etl_ddus.py <файл.xlsx> --pg postgresql://user:pass@host:5432/db

Структура листа: шапка со строкой «Статья оборотов», затем «Итого»,
«Итого текущая деятельность» со статьями, «Итого инвестиционная деятельность»
со статьями. Строки «в т.ч. ...» — детализация предыдущей статьи; они
помечаются is_detail = true и не должны попадать в суммы.
"""
import argparse
import csv
import sys

import openpyxl

SHEET = "Бюджет ДДУС"
# Колонки листа (1-based) -> поле таблицы
VALUE_COLS = {
    3: "sd_5m",            # 5 мес Сд_26
    4: "fact_5m",          # 5 мес Факт_26
    5: "sd_jun_dec",       # Июнь-Дек Сд_26
    6: "forecast_jun_dec", # Июнь-Дек Прогноз_26
    7: "sd_year",          # Сд_26
    8: "expected_year",    # Ожид_26
}
COMMENT_COL = 10
ACTIVITY_HEADERS = {
    "итого текущая деятельность": "Текущая",
    "итого инвестиционная деятельность": "Инвестиционная",
}
FIELDS = ["activity", "article", "parent_article", "is_detail", "sort_order",
          *VALUE_COLS.values(), "comment"]


def num(v):
    if v is None or v == "":
        return 0.0
    return float(v)


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[SHEET]
    rows = list(ws.iter_rows(values_only=True))

    start = next(i for i, r in enumerate(rows) if r[1] == "Статья оборотов")
    activity, parent, out = None, None, []
    for r in rows[start + 1:]:
        name = (r[1] or "").strip() if isinstance(r[1], str) else r[1]
        if not name:
            continue
        key = name.lower()
        if key in ACTIVITY_HEADERS:
            activity = ACTIVITY_HEADERS[key]
            continue
        if activity is None:  # служебные строки шапки и общий «Итого»
            continue
        is_detail = key.startswith("в т.ч.")
        if not is_detail:
            parent = name
        comment = r[COMMENT_COL - 1]
        rec = {
            "activity": activity,
            "article": name,
            "parent_article": parent if is_detail else None,
            "is_detail": is_detail,
            "sort_order": len(out) + 1,
            "comment": comment.strip() if isinstance(comment, str) and comment.strip() else None,
        }
        for col, field in VALUE_COLS.items():
            rec[field] = num(r[col - 1])
        out.append(rec)
    return out


def check_totals(path, recs):
    """Сверка: сумма статей (без «в т.ч.») == строки «Итого» в листе."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = list(wb[SHEET].iter_rows(values_only=True))
    total = next(r for r in rows if r[1] == "Итого")
    ok = True
    for col, field in VALUE_COLS.items():
        s = sum(x[field] for x in recs if not x["is_detail"])
        if abs(s - num(total[col - 1])) > 1:
            print(f"! {field}: статьи {s:,.0f} != Итого {num(total[col - 1]):,.0f}", file=sys.stderr)
            ok = False
    return ok


DDL = """
DROP TABLE IF EXISTS ddus_budget;
CREATE TABLE ddus_budget (
    activity          text,
    article           text,
    parent_article    text,
    is_detail         boolean,
    sort_order        integer,
    sd_5m             double precision,
    fact_5m           double precision,
    sd_jun_dec        double precision,
    forecast_jun_dec  double precision,
    sd_year           double precision,
    expected_year     double precision,
    comment           text
);
"""


def to_pg(recs, url):
    import psycopg2
    with psycopg2.connect(url) as conn, conn.cursor() as cur:
        cur.execute(DDL)
        cols = ", ".join(FIELDS)
        ph = ", ".join(["%s"] * len(FIELDS))
        cur.executemany(f"INSERT INTO ddus_budget ({cols}) VALUES ({ph})",
                        [[r[f] for f in FIELDS] for r in recs])


def to_csv(recs, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(recs)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--csv")
    ap.add_argument("--pg")
    a = ap.parse_args()
    recs = parse(a.xlsx)
    print(f"строк: {len(recs)}; сверка с «Итого»: {'OK' if check_totals(a.xlsx, recs) else 'ОШИБКА'}")
    if a.csv:
        to_csv(recs, a.csv)
    if a.pg:
        to_pg(recs, a.pg)
